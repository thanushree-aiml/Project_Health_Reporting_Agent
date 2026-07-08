from __future__ import annotations

import glob
import json
import os
from collections import Counter
from datetime import datetime
from typing import Any, Dict, List

from config.settings import OUTPUT_FOLDER, REPORT_FOLDER


def _load_inputs(pattern: str) -> List[Dict[str, Any]]:
    paths = sorted(glob.glob(pattern))
    items: List[Dict[str, Any]] = []
    for p in paths:
        with open(p, "r", encoding="utf-8") as f:
            items.append(json.load(f))
    return items


def run_monthly(inputs_pattern: str, output_name: str) -> None:
    os.makedirs(REPORT_FOLDER, exist_ok=True)

    items = _load_inputs(inputs_pattern)
    if not items:
        raise FileNotFoundError(f"No weekly json outputs match pattern: {inputs_pattern}")

    total = len(items)
    overall_counts = Counter([it["rag"]["overall_rag"] for it in items])

    # Trend extraction: list top red/amber projects + compute category trends.
    red_projects = [it["signals"]["project_name"] for it in items if it["rag"]["overall_rag"] == "red"]
    amber_projects = [it["signals"]["project_name"] for it in items if it["rag"]["overall_rag"] == "amber"]

    def category_counts(cat: str) -> Dict[str, int]:
        c = Counter([it["rag"]["per_category"][cat]["rag"] for it in items])
        return dict(c)

    category_trends = {
        "schedule": category_counts("schedule"),
        "budget": category_counts("budget"),
        "milestones": category_counts("milestones"),
        "blockers": category_counts("blockers"),
        "stakeholder_sentiment": category_counts("stakeholder_sentiment"),
    }

    summary = {
        "generated_at": datetime.utcnow().isoformat() + "Z",
        "projects": [it["signals"]["project_name"] for it in items],
        "overall_counts": dict(overall_counts),
        "category_trends": category_trends,
        "top_risks": {
            "red_projects": red_projects,
            "amber_projects": amber_projects,
        },
    }


    # Generate PPTX if python-pptx is available
    try:
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()

        def add_title_slide(title: str, subtitle: str = ""):
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = title
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = subtitle

        def add_bullets_slide(title: str, bullets: List[str]):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = title
            body = slide.shapes.placeholders[1].text_frame
            body.clear()
            for i, b in enumerate(bullets):
                p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                p.text = b
                p.level = 0

        add_title_slide(
            "Monthly Executive Project Health",
            f"Generated {datetime.utcnow().strftime('%Y-%m-%d')} | {total} project(s)",
        )

        add_bullets_slide(
            "Overall trend",
            [
                f"Green: {overall_counts.get('green', 0)}",
                f"Amber: {overall_counts.get('amber', 0)}",
                f"Red: {overall_counts.get('red', 0)}",
                "Focus on Red projects first, then resolve Amber drivers.",
            ],
        )

        # Category trends slide (more executive-friendly than only project lists)
        cat = summary.get("category_trends", {})
        add_bullets_slide(
            "Category distribution (RAG)",
            [
                "Schedule: " + ", ".join([f"{k}={cat.get('schedule', {}).get(k, 0)}" for k in ("green", "amber", "red")]),
                "Budget: " + ", ".join([f"{k}={cat.get('budget', {}).get(k, 0)}" for k in ("green", "amber", "red")]),
                "Milestones: " + ", ".join([f"{k}={cat.get('milestones', {}).get(k, 0)}" for k in ("green", "amber", "red")]),
                "Blockers: " + ", ".join([f"{k}={cat.get('blockers', {}).get(k, 0)}" for k in ("green", "amber", "red")]),
                "Sentiment: " + ", ".join([f"{k}={cat.get('stakeholder_sentiment', {}).get(k, 0)}" for k in ("green", "amber", "red")]),
            ],
        )


        add_bullets_slide(
            "Emerging risks",
            ([
                "Red projects (highest attention):",
                *[f"- {p}" for p in red_projects],
            ] if red_projects else ["Red projects (highest attention):", "- None"]),
        )

        add_bullets_slide(
            "Recommended actions",
            [
                "Re-plan schedule buffers for Red drivers.",
                "Escalate blockers with clear owners + due dates.",
                "Validate budget burn assumptions and procurement timelines.",
            ],
        )

        add_bullets_slide(
            "Assumptions & data quality",
            [
                "RAG is computed from extracted signals; missing critical signals reduce certainty.",
                "Narratives include evidence and confidence per category.",
            ],
        )

        pptx_path = os.path.join(REPORT_FOLDER, output_name)
        prs.save(pptx_path)

    except Exception:
        # If PPTX generation fails, still write JSON summary.
        pptx_path = None

    out_json = os.path.join(REPORT_FOLDER, "monthly_executive_summary.json")
    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2)

    return

